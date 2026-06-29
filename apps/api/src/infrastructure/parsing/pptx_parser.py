import structlog
from pptx import Presentation

from src.infrastructure.parsing.base import AbstractDocumentParser, ParsedDocument, ParsedPage

logger = structlog.get_logger(__name__)


class PptxParser(AbstractDocumentParser):
    @property
    def supported_types(self) -> list[str]:
        return [
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            ".pptx",
        ]

    async def parse(self, file_path: str) -> ParsedDocument:
        try:
            prs = Presentation(file_path)
            pages: list[ParsedPage] = []
            for i, slide in enumerate(prs.slides, start=1):
                texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        texts.append(shape.text.strip())
                content = "\n".join(texts)
                if content:
                    pages.append(ParsedPage(page_number=i, content=content))
            return ParsedDocument(
                pages=pages,
                metadata={"slide_count": len(prs.slides)},
            )
        except Exception as e:
            logger.error("pptx_parse_error", path=file_path, error=str(e))
            raise
