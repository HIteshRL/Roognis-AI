import structlog
from docx import Document as DocxDocument

from src.infrastructure.parsing.base import AbstractDocumentParser, ParsedDocument, ParsedPage

logger = structlog.get_logger(__name__)


class DocxParser(AbstractDocumentParser):
    @property
    def supported_types(self) -> list[str]:
        return [
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ".docx",
        ]

    async def parse(self, file_path: str) -> ParsedDocument:
        try:
            doc = DocxDocument(file_path)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            content = "\n\n".join(paragraphs)
            return ParsedDocument(
                pages=[ParsedPage(page_number=1, content=content)],
                metadata={"paragraph_count": len(paragraphs)},
            )
        except Exception as e:
            logger.error("docx_parse_error", path=file_path, error=str(e))
            raise
